"""Generación del "Informe de Departamento" (Word) a partir de ``informe.docx``.

Se abre la plantilla oficial con ``python-docx`` y se completa:

* encabezado (ciclo lectivo, período, departamento, jefe);
* campos cualitativos (a continuación de cada rótulo subrayado);
* porcentaje de aprobados por año;
* tablas de datos cuantitativos: por asignatura y curso/división, con total de
  alumnos y cantidad / porcentaje de aprobados, desaprobados y ausentes.
"""
from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Twips
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.config import TEMPLATE_PATH
from app.data import repository
from app.data.repository import Departamento
from app.logic import estadisticas
from app.logic.estadisticas import Resumen

# (clave, rótulo tal como aparece en la plantilla)
CAMPOS_CUALITATIVOS: list[tuple[str, str]] = [
    ("desarrollo_programas", "Desarrollo de los Programas:"),
    ("estrategias_metodologicas", "Nivel de Inclusión de las Estrategias Metodológicas:"),
    ("avance_proyectos_6", "Nivel de Avance de los Proyectos de 6º años:"),
    ("clases_apoyo", "Clases de Apoyo y Trayectorias, de Rotación Trimestral:"),
    ("horas_departamentales", "Las horas departamentales son destinadas a:"),
    ("estrategias_dificultades", "Estrategias para el Tratamiento de Dificultades:"),
    ("estado_plan_anual", "Estado de Avance del Plan Anual:"),
    ("logro_objetivos", "Logro de los Objetivos Departamentales:"),
    ("coordinacion", "Coordinación Vertical y Horizontal de los Espacios Curriculares -Integración:"),
    ("aspectos_pos_neg", "Aspectos Positivos y Negativos:"),
    ("criterios_evaluacion", "Criterios de Evaluación:"),
    ("perfeccionamiento", "Participación en Actividades de Perfeccionamiento y Actualización Docente:"),
    ("sugerencias", "Sugerencias:"),
]

# Anchos de columna (twips) de la tabla del ciclo básico. Las tablas del ciclo
# orientado de la plantilla traen otra grilla; se normalizan a esta para que
# todas las columnas queden alineadas.
_ANCHOS = (2518, 1276, 1154, 547, 1134, 566, 1135, 596, 963)

_RE_CURSO = re.compile(r"^\s*(\d)\s*[°º]\s*(\d+)\s*[°º]\s*$")


@dataclass
class DatosInforme:
    """Datos que el usuario aporta para el informe."""

    departamento: Departamento
    ciclo_lectivo: str
    periodo: str
    textos: dict[str, str] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Utilidades de edición sobre texto con runs
# ---------------------------------------------------------------------------
def _reemplazar(par: Paragraph, patron: str, nuevo: str) -> bool:
    """Reemplaza la primera coincidencia de ``patron`` aunque cruce varios runs.

    Conserva el formato del primer run afectado.
    """
    m = re.search(patron, par.text)
    if not m:
        return False
    ini, fin = m.span()
    pos, primero = 0, True
    for run in par.runs:
        r_ini, r_fin = pos, pos + len(run.text)
        pos = r_fin
        if r_fin <= ini or r_ini >= fin:
            continue
        a, b = max(ini, r_ini) - r_ini, min(fin, r_fin) - r_ini
        t = run.text
        run.text = t[:a] + (nuevo if primero else "") + t[b:]
        primero = False
    return True


def _fmt_pct(valor: float) -> str:
    return f"{valor:.1f}".replace(".", ",") + " %"


def _completar_rotulo(par: Paragraph, rotulo: str, valor: str) -> None:
    """Escribe ``valor`` a continuación del rótulo, descartando lo que hubiera
    después (p. ej. las rayas ``-----`` de la plantilla)."""
    patron_rotulo = r"\s+".join(re.escape(w) for w in rotulo.split())
    m = re.search(patron_rotulo, par.text)
    if not m:
        return
    fin = m.end()
    pos = 0
    ultimo_con_formato = None
    for run in par.runs:
        r_ini, r_fin = pos, pos + len(run.text)
        pos = r_fin
        if r_ini >= fin:
            run.text = ""
        elif r_fin > fin:
            run.text = run.text[: fin - r_ini]
        if run.text:
            ultimo_con_formato = run
    nuevo = par.add_run(" " + valor.strip())
    nuevo.bold = nuevo.italic = nuevo.underline = False
    if ultimo_con_formato is not None:
        nuevo.font.name = ultimo_con_formato.font.name
        nuevo.font.size = ultimo_con_formato.font.size


def _buscar_parrafo(doc, rotulo: str) -> Optional[Paragraph]:
    norm = " ".join(rotulo.split())
    for p in doc.paragraphs:
        if " ".join(p.text.split()).startswith(norm):
            return p
    return None


# ---------------------------------------------------------------------------
# Tablas
# ---------------------------------------------------------------------------
def _es_fila_datos(tr, tabla: Table) -> bool:
    tcs = tr.tc_lst
    return len(tcs) == 9 and bool(_RE_CURSO.match(_Cell(tcs[1], tabla).text))


def _escribir_celda(tc, tabla: Table, texto: str, rpr) -> None:
    """Deja ``texto`` como único contenido de la celda, con formato ``rpr``."""
    par = _Cell(tc, tabla).paragraphs[0]
    for r in list(par._p.r_lst):
        par._p.remove(r)
    run = par.add_run(texto)
    if rpr is not None:
        viejo = run._r.rPr
        if viejo is not None:
            run._r.remove(viejo)
        run._r.insert(0, deepcopy(rpr))


def _reconstruir_tabla(tabla: Table, grupos: dict[tuple[str, int, int], Resumen]) -> None:
    """Reemplaza las filas de ejemplo de la plantilla por filas con datos reales."""
    trs = [tr for tr in tabla._tbl.tr_lst if _es_fila_datos(tr, tabla)]
    if not trs:
        return
    anios = {int(_RE_CURSO.match(_Cell(tr.tc_lst[1], tabla).text).group(1)) for tr in trs}

    proto = deepcopy(trs[0])
    runs_proto = proto.tc_lst[1].xpath(".//w:r")
    rpr = runs_proto[0].rPr if runs_proto else None

    padre = trs[0].getparent()
    indice = padre.index(trs[0])
    for tr in trs:
        padre.remove(tr)

    # Orden: curso, división y asignatura (como en la plantilla, por curso/división).
    filas = []
    for (materia, anio, division), r in sorted(
            grupos.items(), key=lambda kv: (kv[0][1], kv[0][2], kv[0][0])):
        if anio not in anios or r.total == 0:
            continue
        sep = "°" if anio <= 3 else "º"
        filas.append([materia, f"{anio}{sep}{division}{sep}", str(r.total),
                      str(r.aprobados), _fmt_pct(r.pct_aprobados),
                      str(r.desaprobados), _fmt_pct(r.pct_desaprobados),
                      str(r.ausentes), _fmt_pct(r.pct_ausentes)])
    if not filas:                                       # ej. "NO HAY" en ciclo básico
        filas = [["NO HAY", "-", "-", "-", "-", "-", "-", "-", "-"]]

    for i, valores in enumerate(filas):
        tr = deepcopy(proto)
        for tc, texto, ancho in zip(tr.tc_lst, valores, _ANCHOS):
            _escribir_celda(tc, tabla, texto, rpr)
            tc.width = Twips(ancho)
        padre.insert(indice + i, tr)

    for col, ancho in zip(tabla._tbl.tblGrid.gridCol_lst, _ANCHOS):
        col.w = Twips(ancho)


def _completar_porcentajes_por_anio(tabla: Table, por_anio: dict[int, Resumen]) -> None:
    """Tabla resumen "Porcentaje de aprobados" (1º a 6º año)."""
    for fila in tabla.rows:
        for tc in fila._tr.tc_lst:
            celda = _Cell(tc, tabla)
            m = re.match(r"\s*(\d)\s*[º°]\s*a[ñn]o", celda.text)
            if not m:
                continue
            r = por_anio.get(int(m.group(1)))
            valor = _fmt_pct(r.pct_aprobados) if r and r.total else "-"
            _reemplazar(celda.paragraphs[0], r"(?<=año:).*", " " + valor)


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------
def generar_informe(datos: DatosInforme, ruta_salida: Path | str,
                    plantilla: Path | str = TEMPLATE_PATH) -> Path:
    """Genera el informe del departamento y lo guarda en ``ruta_salida``.

    Consulta la base de datos, evalúa cada alumno con las reglas de negocio
    (``app.logic``) y completa la plantilla ``informe.docx``.
    """
    plantilla = Path(plantilla)
    if not plantilla.exists():
        raise FileNotFoundError(f"No se encontró la plantilla: {plantilla}")

    filas = repository.listar_filas(departamento_id=datos.departamento.id)
    grupos = estadisticas.agrupar_por_materia_y_curso(filas)
    por_anio = estadisticas.agrupar_por_anio(grupos)

    doc = Document(str(plantilla))

    # --- Encabezado --------------------------------------------------------
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("CICLO LECTIVO"):
            _reemplazar(p, r"\d{4}", str(datos.ciclo_lectivo))
        elif "EVALUACIÓN Y SEGUIMIENTO" in t:
            _reemplazar(p, r"[….]+\s*TRIMESTRE", datos.periodo.upper())
        elif t.startswith("DEPARTAMENTO DE"):
            _reemplazar(p, r"(?<=DEPARTAMENTO DE ).*", datos.departamento.nombre.upper())
        elif t.startswith("JEFE DE"):
            _reemplazar(p, r"(?<=DEPARTAMENTO:).*", " " + datos.departamento.jefe.upper())

    # --- Informe cualitativo -----------------------------------------------
    for clave, rotulo in CAMPOS_CUALITATIVOS:
        valor = (datos.textos.get(clave) or "").strip()
        if not valor:
            continue
        par = _buscar_parrafo(doc, rotulo)
        if par is not None:
            _completar_rotulo(par, rotulo, valor)

    # --- Datos cuantitativos -----------------------------------------------
    for tabla in doc.tables:
        if any("año:" in c.text for c in tabla.rows[0].cells) or \
           "Ciclo Básico:" in tabla.rows[0].cells[0].text:
            _completar_porcentajes_por_anio(tabla, por_anio)
        else:
            _reconstruir_tabla(tabla, grupos)

    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ruta_salida))
    return ruta_salida
